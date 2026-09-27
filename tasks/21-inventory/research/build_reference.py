"""Build the task-21 reference answer as a LIVE-FORMULA workbook (src/pips-q3-report.xlsx): an independent method from
engine.py (cumulative sums instead of a FIFO queue), recalculated by LibreOffice. Swapping the Transactions sheet must
keep it correct — that is exactly what the grader's live test does.
    python tasks/21-inventory/research/build_reference.py
"""
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parents[1]
N = 1001                     # Transactions rows 2..N (capacity 1,000 transactions)
MONTHS = [202507, 202508, 202509]
LABELS = ["2025-07", "2025-08", "2025-09"]

wb = load_workbook(ROOT / "inputs" / "pips-q3.xlsx")
T = "Transactions"
c = wb.create_sheet("Calc")
heads = ["SKU", "Type", "Month", "Dup", "Valid", "SupplyQty", "SupplyCost", "RcptRow", "DemandQty", "SupplyBefore", "DemandBefore",
         "OpenQty", "Ship1", "Ship2", "Ship3", "C0", "C1", "C2", "C3", "Refund", "Opex", "ExcRow", "OpenCost"]
c.append(heads)
for cell in c[1]: cell.font = Font(bold=True)
rng = lambda col: f"${col}$2:${col}${N}"


def cost_fn(n, sku_ref, L, oc):
    """FIFO cost of the first n units of SKU: the opening layer first, then each supply row's layer in order."""
    x = f"({n}-({L}+Calc!{rng('J')}))"
    return (f"(MIN(MAX({n},0),{L})*{oc}+SUMPRODUCT((Calc!{rng('A')}={sku_ref})*Calc!{rng('G')}*"
            f"(({x}>0)*({x}<Calc!{rng('F')})*{x}+({x}>=Calc!{rng('F')})*Calc!{rng('F')})))")


for r in range(2, N + 1):
    p = r - 1
    t = lambda col: f"{T}!{col}{r}"
    mrow = f"MATCH($A{r},Month!$A$2:$A$13,0)"
    row = [
        f'=IF({t("D")}="","",UPPER(TRIM({t("D")})))',
        f'=IF({t("C")}="","",UPPER(TRIM({t("C")})))',
        f'=IF({t("B")}="",0,YEAR({t("B")})*100+MONTH({t("B")}))',
        f'=IF({t("A")}="",1,IF(COUNTIF({T}!$A$1:A{p},{t("A")})>0,1,0))',
        f'=IF(D{r}=1,0,IF(B{r}="RETURN",IF(COUNTIFS($A$1:A{p},A{r},$B$1:B{p},"SALE",$E$1:E{p},1)>0,1,0),1))',
        f'=IF(E{r}=1,IF(B{r}="RECEIPT",{t("E")},IF(AND(B{r}="RETURN",LOWER(TRIM({t("H")}))="resellable"),{t("E")},0)),0)',
        f'=IF(F{r}>0,IF(B{r}="RECEIPT",{t("F")},IF(_xlfn.MAXIFS($H$2:H{r},$A$2:A{r},A{r})>0,INDEX({T}!$F:$F,_xlfn.MAXIFS($H$2:H{r},$A$2:A{r},A{r})),W{r})),0)',
        f'=IF(AND(E{r}=1,B{r}="RECEIPT"),ROW(),0)',
        f'=IF(E{r}=1,IF(OR(B{r}="SALE",B{r}="WRITEOFF"),{t("E")},0),0)',
        f'=SUMIFS($F$1:F{p},$A$1:A{p},A{r})',
        f'=SUMIFS($I$1:I{p},$A$1:A{p},A{r})',
        f'=IFERROR(VLOOKUP(A{r},Opening!$A:$C,3,0),0)',
    ]
    for k in range(3):  # units of this demand row shipped by the end of month k
        row.append(f'=IF(I{r}=0,0,MAX(0,MIN(INDEX(Month!{"HIJ"[k]}$2:{"HIJ"[k]}$13,{mrow})-K{r},I{r})))')
    L, oc = f"L{r}", f"W{r}"
    row.append(f'=IF(I{r}=0,0,{cost_fn(f"K{r}", f"A{r}", L, oc)})')
    for k in range(3):
        row.append(f'=IF(I{r}=0,0,{cost_fn(f"(K{r}+{"MNO"[k]}{r})", f"A{r}", L, oc)})')
    row += [f'=IF(AND(E{r}=1,B{r}="RETURN"),{t("E")}*{t("G")},0)',
            f'=IF(AND(E{r}=1,B{r}="EXPENSE"),{t("J")},0)',
            f'=IF(AND({t("A")}<>"",E{r}=0),ROW(),"")',
            f'=IFERROR(VLOOKUP(A{r},Opening!$A:$D,4,0),0)']
    c.append(row)

m = wb.create_sheet("Month")
m.append(["SKU", "Demand≤Jul", "Demand≤Aug", "Demand≤Sep", "Supply≤Jul", "Supply≤Aug", "Supply≤Sep", "Shipped≤Jul", "Shipped≤Aug", "Shipped≤Sep"])
for i in range(2, 14):
    row = [f"=Opening!A{i}"]
    row += [f'=SUMIFS(Calc!{rng("I")},Calc!{rng("A")},$A{i},Calc!{rng("C")},"<="&{mm})' for mm in MONTHS]
    row += [f'=Opening!C{i}+SUMIFS(Calc!{rng("F")},Calc!{rng("A")},$A{i},Calc!{rng("C")},"<="&{mm})' for mm in MONTHS]
    row += [f"=MIN({'BCD'[k]}{i},{'EFG'[k]}{i})" for k in range(3)]
    m.append(row)

s = wb.create_sheet("Summary")
s.append(["Pip's Outdoor Gear — Q3 2025 P&L"])
s.append(["Metric"] + LABELS + ["Q3 total"])
metrics = ["Revenue", "Refunds", "COGS", "Write-offs", "Gross profit", "Operating expenses", "Net profit"]
for j, name in enumerate(metrics):
    r = j + 3
    row = [name]
    for k, mm in enumerate(MONTHS):
        prev_ship = "0" if k == 0 else f'Calc!{rng("MNO"[k - 1])}'
        prev_c = f'Calc!{rng("P")}' if k == 0 else f'Calc!{rng("QRS"[k - 1])}'
        ship, cc = f'Calc!{rng("MNO"[k])}', f'Calc!{rng("QRS"[k])}'
        col = "BCD"[k]
        if name == "Revenue": f = f'=SUMPRODUCT((Calc!{rng("B")}="SALE")*{T}!{rng("G")}*({ship}-{prev_ship}))'
        elif name == "Refunds": f = f'=SUMIFS(Calc!{rng("T")},Calc!{rng("C")},{mm})'
        elif name == "COGS": f = f'=SUMPRODUCT((Calc!{rng("B")}="SALE")*({cc}-{prev_c}))'
        elif name == "Write-offs": f = f'=SUMPRODUCT((Calc!{rng("B")}="WRITEOFF")*({cc}-{prev_c}))'
        elif name == "Gross profit": f = f"={col}3-{col}4-{col}5-{col}6"
        elif name == "Operating expenses": f = f'=SUMIFS(Calc!{rng("U")},Calc!{rng("C")},{mm})'
        else: f = f"={col}7-{col}8"
        row.append(f)
    row.append(f"=SUM(B{r}:D{r})")
    s.append(row)

st = wb.create_sheet("Stock")
st.append(["SKU", "On hand", "Backordered", "Stock value (FIFO)", "Reorder?", "Reorder qty"])
for i in range(2, 14):
    L, oc = f"Opening!C{i}", f"Opening!D{i}"
    supply, shipped = f"Month!G{i}", f"Month!J{i}"
    st.append([f"=Opening!A{i}", f"={supply}-{shipped}",
               f'=SUMPRODUCT((Calc!{rng("A")}=$A{i})*(Calc!{rng("B")}="SALE")*(Calc!{rng("I")}-Calc!{rng("O")}))',
               f"=ROUND({cost_fn(supply, f'$A{i}', L, oc)}-{cost_fn(shipped, f'$A{i}', L, oc)},2)",
               f'=IF(B{i}<Opening!F{i},"YES","NO")', f"=IF(E{i}=\"YES\",Opening!G{i},0)"])
ex = wb.create_sheet("Exceptions")
ex.append(["TxnID", "Why it was ignored"])
for k in range(1, 41):
    r = k + 1
    ex.append([f'=IFERROR(INDEX({T}!$A:$A,SMALL(Calc!{rng("V")},{k})),"")',
               f'=IF(A{r}="","",IF(COUNTIF({T}!$A$1:INDEX({T}!$A:$A,SMALL(Calc!{rng("V")},{k})-1),A{r})>0,"duplicate transaction id","return of a product never sold"))'])
for sh in (m, s, st, ex):
    for cell in sh[1]: cell.font = Font(bold=True)
out = ROOT / "src" / "pips-q3-report.xlsx"
out.parent.mkdir(exist_ok=True)
wb.save(out)
print("wrote", out)
