"""Task 21 grader core: recalc the model's workbook in LibreOffice, score its report against the answer key, then swap in
the hidden transaction log (the live test), recalc again and score again. Prints {checks, notes} JSON.
    python grade_core.py <path/to/pips-q3-report.xlsx>
"""
import json, sys, traceback
from pathlib import Path
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from recalc import recalc          # noqa: E402
from readout import read_report, MONTHS, METRICS  # noqa: E402
from swap import swap_transactions  # noqa: E402

TOL = 0.011
checks, notes = [], []
add = lambda id, group, name, points, earned, detail: checks.append(dict(id=id, group=group, name=name, points=points, earned=earned, detail=detail))
num = lambda v: isinstance(v, (int, float)) and not isinstance(v, bool)
close = lambda g, e: num(g) and abs(g - e) <= TOL
ERRS = ("#REF!", "#VALUE!", "#NAME?", "#DIV/0!", "#N/A", "#NUM!", "Err:")


def report_errors(path):
    wb = load_workbook(path, data_only=True); bad = []
    for sh, cells in (("Summary", "B3:E9"), ("Stock", "B2:F13")):
        if sh not in wb.sheetnames: continue
        for row in wb[sh][cells]:
            for c in row:
                if isinstance(c.value, str) and c.value.startswith(ERRS): bad.append(f"{sh}!{c.coordinate}={c.value}")
    return bad


def score_numbers(got, key):
    """[(label, ok)] for the 21 monthly P&L cells + 3×12 stock numbers."""
    out = []
    for m in MONTHS:
        for k in METRICS: out.append((f"{m} {k}", close(got["pl"][m].get(k), key["pl"][m][k])))
    for sku, v in key["inventory"].items():
        g = got["stock"].get(sku, {})
        for k in ("on_hand", "backordered", "stock_value"): out.append((f"{sku} {k}", close(g.get(k), v[k])))
    return out


def main(path):
    path = Path(path)
    key = json.loads((HERE / "answer-given.json").read_text()); hkey = json.loads((HERE / "answer-hidden.json").read_text())
    if not path.exists():
        notes.append("no src/pips-q3-report.xlsx"); return
    try: wb0 = load_workbook(path)
    except Exception as e:
        notes.append(f"workbook does not open: {e}"); return
    add("F-opens", "Deliverable", "workbook opens", 1, 1, "ok")
    need = [s for s in ("Summary", "Stock", "Exceptions") if s not in wb0.sheetnames]
    add("F-sheets", "Deliverable", "Summary, Stock and Exceptions sheets exist", 1, 0 if need else 1, f"missing {need}" if need else "all present")
    tx_ok = "Transactions" in wb0.sheetnames and [c.value for c in wb0["Transactions"][1][:11]] == ["TxnID", "Date", "Type", "SKU", "Qty", "UnitCost", "UnitPrice", "Condition", "Category", "Amount", "Note"] and "Opening" in wb0.sheetnames
    add("F-inputs", "Deliverable", "Opening and Transactions sheets kept with the same columns", 1, 1 if tx_ok else 0, "ok" if tx_ok else "Transactions/Opening sheet missing or its header changed")
    # Formula discipline: report cells must be formulas (they are what makes the workbook live).
    fcells = fn = 0
    for sh, cells in (("Summary", "B3:E9"), ("Stock", "B2:F13")):
        if sh in wb0.sheetnames:
            for row in wb0[sh][cells]:
                for c in row: fcells += 1; fn += isinstance(c.value, str) and c.value.startswith("=")
    add("F-formulas", "Deliverable", "report cells are formulas, not typed-in numbers", 2, 2 if fcells and fn == fcells else 1 if fcells and fn >= 0.8 * fcells else 0, f"{fn}/{fcells} report cells are formulas")

    rec = recalc(path)
    errs = report_errors(rec)
    add("F-errors", "Deliverable", "no formula errors in the report after recalculation", 1, 0 if errs else 1, ", ".join(errs[:4]) or "none")
    got = read_report(rec)
    pl, K = got["pl"], key["pl"]
    for m in MONTHS:
        add(f"P-rev-{m}", "P&L (this quarter)", f"revenue {m} (shipped units, incl. backorders)", 1, 1 if close(pl[m].get("revenue"), K[m]["revenue"]) else 0, f"got {pl[m].get('revenue')} want {K[m]['revenue']}")
        add(f"P-cogs-{m}", "P&L (this quarter)", f"COGS {m} (FIFO)", 1, 1 if close(pl[m].get("cogs"), K[m]["cogs"]) else 0, f"got {pl[m].get('cogs')} want {K[m]['cogs']}")
        add(f"P-net-{m}", "P&L (this quarter)", f"net profit {m}", 1, 1 if close(pl[m].get("net_profit"), K[m]["net_profit"]) else 0, f"got {pl[m].get('net_profit')} want {K[m]['net_profit']}")
    for k, name in (("refunds", "refunds"), ("writeoffs", "write-offs"), ("opex", "operating expenses (duplicates dropped)"), ("gross_profit", "gross profit")):
        ok = all(close(pl[m].get(k), K[m][k]) for m in MONTHS)
        add(f"P-{k}", "P&L (this quarter)", f"{name}, all three months", 1, 1 if ok else 0, "; ".join(f"{m}: got {pl[m].get(k)} want {K[m][k]}" for m in MONTHS if not close(pl[m].get(k), K[m][k])) or "all match")
    tot_ok = all(close(got["totals"].get(k), key["totals"][k]) for k in METRICS)
    add("P-totals", "P&L (this quarter)", "Q3 total column", 1, 1 if tot_ok else 0, "all match" if tot_ok else "; ".join(f"{k}: got {got['totals'].get(k)} want {key['totals'][k]}" for k in METRICS if not close(got["totals"].get(k), key["totals"][k]))[:300])
    inv = key["inventory"]
    for k, pts, name in (("on_hand", 2, "units on hand"), ("backordered", 2, "units still backordered"), ("stock_value", 3, "FIFO stock value")):
        good = [s for s, v in inv.items() if close(got["stock"].get(s, {}).get(k), v[k])]
        earned = pts if len(good) == 12 else (pts - 1 if len(good) >= 10 else (1 if len(good) >= 6 and pts == 3 else 0))
        miss = [f"{s}: got {got['stock'].get(s, {}).get(k)} want {v[k]}" for s, v in inv.items() if s not in good]
        add(f"S-{k}", "Stock (this quarter)", f"{name}, all 12 products", pts, earned, f"{len(good)}/12 right" + (f" — {'; '.join(miss[:3])}" if miss else ""))
    ro = [s for s, v in inv.items() if (got["stock"].get(s, {}).get("reorder") == "YES") == v["reorder"] and close(got["stock"].get(s, {}).get("reorder_qty") or 0, v["reorder_qty"])]
    add("S-reorder", "Stock (this quarter)", "reorder flags and quantities", 2, 2 if len(ro) == 12 else 1 if len(ro) >= 10 else 0, f"{len(ro)}/12 right")
    want = {e[0] for e in key["exceptions"]}; listed = set(got["exceptions"])
    add("E-found", "Exceptions", "every ignored row is listed (duplicates + invalid return)", 2, 2 if want <= listed else 1 if want & listed else 0, f"want {sorted(want)}; listed {sorted(listed)[:8]}")
    add("E-clean", "Exceptions", "no valid rows wrongly listed as ignored", 1, 1 if listed and listed <= want else 0, f"extra: {sorted(listed - want)[:6]}" if listed - want else "none extra")

    # ---------------- the live test: swap in the hidden log, recalc, score again
    hidden = json.loads((HERE / "hidden-transactions.json").read_text())
    try:
        swapped = swap_transactions(path, hidden, rec.parent / "swapped.xlsx")
        hrec = recalc(swapped)
        herrs = report_errors(hrec)
        hgot = read_report(hrec)
        res = score_numbers(hgot, hkey); ok = sum(o for _, o in res); frac = ok / len(res)
        add("L-live", "Live test (hidden log)", "with next quarter's log pasted in, the report is still right", 12,
            12 if frac == 1 else 9 if frac >= 0.9 else 6 if frac >= 0.7 else 3 if frac >= 0.4 else 0,
            f"{ok}/{len(res)} numbers right after swapping the log" + (f" — wrong: {', '.join(l for l, o in res if not o)[:220]}" if ok < len(res) else ""))
        add("L-errors", "Live test (hidden log)", "no formula errors after the swap", 1, 0 if herrs else 1, ", ".join(herrs[:4]) or "none")
    except Exception as e:
        notes.append(f"live test failed to run: {str(e)[:200]}")


if __name__ == "__main__":
    try: main(sys.argv[1] if len(sys.argv) > 1 else str(HERE.parent / "src" / "pips-q3-report.xlsx"))
    except Exception: notes.append("grader error: " + traceback.format_exc()[-400:])
    print(json.dumps({"checks": checks, "notes": notes}))
