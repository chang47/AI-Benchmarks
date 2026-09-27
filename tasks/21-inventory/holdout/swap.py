"""The live test: replace every data row of the workbook's Transactions sheet with another log (same columns), keeping
all formulas elsewhere untouched, and save a copy for recalculation."""
import datetime as dt, json
from pathlib import Path
from openpyxl import load_workbook

COLS = ["id", "date", "type", "sku", "qty", "unit_cost", "unit_price", "condition", "category", "amount", "note"]


def swap_transactions(src, rows, dst):
    wb = load_workbook(src)  # formulas preserved (not data_only)
    if "Transactions" not in wb.sheetnames: raise ValueError("no Transactions sheet")
    tx = wb["Transactions"]
    for r in range(2, max(tx.max_row, 2) + 1):
        for c in range(1, 12): tx.cell(r, c).value = None
    for i, row in enumerate(rows, start=2):
        for c, k in enumerate(COLS, start=1):
            v = row.get(k)
            if k == "date" and isinstance(v, str): v = dt.date.fromisoformat(v)
            tx.cell(i, c).value = v
            if k == "date": tx.cell(i, c).number_format = "yyyy-mm-dd"
    wb.save(dst)
    return Path(dst)
