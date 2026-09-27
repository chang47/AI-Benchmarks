"""Read the fixed report layout (Summary + Stock + Exceptions) from a recalculated workbook."""
from openpyxl import load_workbook

METRICS = ["revenue", "refunds", "cogs", "writeoffs", "gross_profit", "opex", "net_profit"]
MONTHS = ["2025-07", "2025-08", "2025-09"]


def read_report(path):
    wb = load_workbook(path, data_only=True)
    out = {"errors": [], "pl": {m: {} for m in MONTHS}, "totals": {}, "stock": {}, "exceptions": []}
    for need in ("Summary", "Stock", "Exceptions"):
        if need not in wb.sheetnames: out["errors"].append(f"missing sheet {need}")
    if "Summary" in wb.sheetnames:
        s = wb["Summary"]
        for j, key in enumerate(METRICS):
            for k, m in enumerate(MONTHS): out["pl"][m][key] = s.cell(3 + j, 2 + k).value
            out["totals"][key] = s.cell(3 + j, 5).value
    if "Stock" in wb.sheetnames:
        st = wb["Stock"]
        for i in range(2, 14):
            sku = str(st.cell(i, 1).value or "").strip().upper()
            out["stock"][sku] = {"on_hand": st.cell(i, 2).value, "backordered": st.cell(i, 3).value, "stock_value": st.cell(i, 4).value,
                                 "reorder": str(st.cell(i, 5).value or "").strip().upper(), "reorder_qty": st.cell(i, 6).value}
    if "Exceptions" in wb.sheetnames:
        ex = wb["Exceptions"]
        out["exceptions"] = [str(ex.cell(r, 1).value).strip() for r in range(2, ex.max_row + 1) if ex.cell(r, 1).value not in (None, "")]
    return out
